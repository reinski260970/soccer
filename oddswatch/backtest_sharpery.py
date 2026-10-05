"""OOS Sharpery Backtest: fester Entry-Filter auf unangetastetem Holdout 2025/26.

- Modellvariante + Kalibrierung werden ausschließlich auf Tune 2023/24+2024/25 gewählt.
- Holdout-Regel wird NICHT optimiert: Modell-EV >= 3%, Opening Odds <= 6.00, alle 1X2-Seiten.
- Jede historische Holdout-Wette wird als Sharpery-Zeile exportiert.
- 1 EH = 10 EUR.
"""

from __future__ import annotations

import csv
import math
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from . import pricing
from .m13_research import LEAGUES as M13_LEAGUES, VARIANTS as M13_VARIANTS, CALIB as M13_CALIB
from .m13_research import load_real_xg, model_score as m13_score
from .m14_research import LEAGUES as M14_LEAGUES, VARIANTS as M14_VARIANTS, CALIB as M14_CALIB
from .m14_research import load_league, train_proxy
from .models.m11 import M11Model
from .models.m9 import M9Model

COLUMNS = ["Placed At", "Kickoff", "Event", "Sport", "League", "Market", "Selection",
           "Period", "Bookmaker", "Opening Odds", "Bet Odds", "Closing Novig Odds",
           "EV %", "CLV %", "Stake", "Result", "Liquidity", "Live", "Side", "Notes", "Tags"]

SIDE = ("home", "draw", "away")
SEL = ("Home", "Draw", "Away")
MIN_EDGE = 0.03
MAX_ODDS = 6.0
STAKE_EUR = 10.0


def _calibrate(p: list[float], a: float) -> list[float]:
    q = [max(x, 1e-9) ** a for x in p]
    z = sum(q)
    return [x / z for x in q]


def _pick_model(matches, rows, variants, calibs, model_cls):
    tune = {2023, 2024}
    best = None
    for vi, params in enumerate(variants):
        out = []
        model = None
        last_fit = None
        for season, d, h, a, hg, ag, op, cl in rows:
            if season not in tune:
                continue
            if last_fit is None or (d - last_fit).days >= 14:
                hist = [m for m in matches if m.date < d]
                try:
                    model = model_cls.fit(hist, d, params)
                except ValueError:
                    model = None
                last_fit = d
            if model is None:
                continue
            try:
                mk = model.markets(h, a, kickoff=d)
            except Exception:
                continue
            pm = [mk["1"], mk["X"], mk["2"]]
            out.append((season, pm, pricing.devig(op), pricing.devig(cl),
                        [hg > ag, hg == ag, hg < ag], op))
        if len(out) < 300:
            continue
        for a in calibs:
            lm = 0.0
            for _, pm, _, _, y, _ in out:
                p = _calibrate(pm, a)
                i = y.index(True)
                lm -= math.log(max(p[i], 1e-12))
            ll = lm / len(out)
            if best is None or ll < best["logloss"]:
                best = {"variant": vi, "params": params, "calib": a,
                        "logloss": ll, "n": len(out)}
    return best


def _holdout_rows(matches, rows, model_cls, params, calib, league, model_tag):
    out = []
    model = None
    last_fit = None
    for season, d, h, a, hg, ag, op, cl in rows:
        if season != 2025:
            continue
        if last_fit is None or (d - last_fit).days >= 14:
            hist = [m for m in matches if m.date < d]
            try:
                model = model_cls.fit(hist, d, params)
            except ValueError:
                model = None
            last_fit = d
        if model is None:
            continue
        try:
            mk = model.markets(h, a, kickoff=d)
        except Exception:
            continue
        p = _calibrate([mk["1"], mk["X"], mk["2"]], calib)
        close_p = pricing.devig(cl)
        result_idx = 0 if hg > ag else 1 if hg == ag else 2
        for k in range(3):
            ev = p[k] * op[k] - 1.0
            if ev < MIN_EDGE or op[k] > MAX_ODDS:
                continue
            clv = op[k] * close_p[k] - 1.0
            pnl = op[k] - 1.0 if k == result_idx else -1.0
            placed = datetime.combine(d, datetime.min.time(), tzinfo=timezone.utc).isoformat()
            fair = 1.0 / max(p[k], 1e-12)
            close_fair = 1.0 / max(close_p[k], 1e-12)
            out.append({
                "Placed At": placed,
                "Kickoff": placed,
                "Event": f"{h} vs {a}",
                "Sport": "Soccer",
                "League": league,
                "Market": "1X2",
                "Selection": SEL[k],
                "Period": "Regular Time",
                "Bookmaker": "Pinnacle",
                "Opening Odds": round(op[k], 4),
                "Bet Odds": round(op[k], 4),
                "Closing Novig Odds": round(close_fair, 4),
                "EV %": round(ev * 100, 2),
                "CLV %": round(clv * 100, 2),
                "Stake": STAKE_EUR,
                "Result": "won" if k == result_idx else "lost",
                "Liquidity": None,
                "Live": "no",
                "Side": SIDE[k],
                "Notes": f"OOS Holdout 2025/26; {model_tag}; Fair {fair:.3f}; fixed edge >=3%; no holdout tuning",
                "Tags": f"BACKTEST;OOS;CLV;{model_tag}",
                "_pnl_eh": pnl,
            })
    return out


def run(out_base: str = "exports/sharpery_oos_backtest"):
    bets = []
    diagnostics = []

    for league, code in M13_LEAGUES.items():
        ms, rows, cov = load_real_xg(code)
        best = _pick_model(ms, rows, M13_VARIANTS, M13_CALIB, M11Model)
        if best is None:
            diagnostics.append((league, "M13", "no tune model"))
            continue
        rr = _holdout_rows(ms, rows, M11Model, best["params"], best["calib"],
                           league, f"M13-v{best['variant']}-a{best['calib']:.2f}")
        bets.extend(rr)
        diagnostics.append((league, "M13", len(rr)))

    proxy = train_proxy()
    for league, code in M14_LEAGUES.items():
        ms, rows, cov = load_league(code, proxy)
        best = _pick_model(ms, rows, M14_VARIANTS, M14_CALIB, M9Model)
        if best is None:
            diagnostics.append((league, "M14", "no tune model"))
            continue
        rr = _holdout_rows(ms, rows, M9Model, best["params"], best["calib"],
                           league, f"M14-v{best['variant']}-a{best['calib']:.2f}")
        bets.extend(rr)
        diagnostics.append((league, "M14", len(rr)))

    bets.sort(key=lambda r: (r["Kickoff"], r["League"], r["Event"], r["Side"]))
    base = Path(out_base)
    base.parent.mkdir(parents=True, exist_ok=True)
    csvp, xlsxp = base.with_suffix(".csv"), base.with_suffix(".xlsx")

    with csvp.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        for r in bets:
            w.writerow({k: "" if r.get(k) is None else r.get(k) for k in COLUMNS})

    wb = Workbook()
    ws = wb.active
    ws.title = "Bets"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in bets:
        ws.append([r.get(k) for k in COLUMNS])
    ws.freeze_panes = "A2"

    sm = wb.create_sheet("Summary")
    sm.append(["Metric", "Value"])
    sm["A1"].font = sm["B1"].font = Font(bold=True)
    n = len(bets)
    avg_clv = sum(r["CLV %"] for r in bets) / n if n else 0.0
    pos_clv = sum(1 for r in bets if r["CLV %"] > 0) / n if n else 0.0
    pnl_eh = sum(r["_pnl_eh"] for r in bets)
    sm.append(["Bets", n])
    sm.append(["Average CLV %", round(avg_clv, 3)])
    sm.append(["Positive CLV rate %", round(pos_clv * 100, 2)])
    sm.append(["ROI %", round((pnl_eh / n * 100) if n else 0.0, 2)])
    sm.append(["Stake EUR", round(n * STAKE_EUR, 2)])
    sm.append(["P/L EUR", round(pnl_eh * STAKE_EUR, 2)])
    sm.append(["Rule", "OOS 2025/26; fixed model EV >=3%; max odds 6.00; 1 EH=10 EUR"])
    for lg, model, count in diagnostics:
        sm.append([f"{lg} {model}", count])

    for wsx in (ws, sm):
        for col in wsx.columns:
            width = max(len(str(c.value or "")) for c in col)
            wsx.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 60)
    wb.save(xlsxp)

    return xlsxp, csvp, bets, diagnostics


def main():
    x, c, bets, diag = run()
    n = len(bets)
    avg = sum(r["CLV %"] for r in bets)/n if n else 0.0
    pos = sum(1 for r in bets if r["CLV %"] > 0)/n if n else 0.0
    pnl = sum(r["_pnl_eh"] for r in bets)
    print(f"OOS BETS={n} AVG_CLV={avg:+.3f}% POS_CLV={pos*100:.1f}% ROI={(pnl/n*100 if n else 0):+.2f}%")
    print("DIAG", diag)
    print(x)
    print(c)


if __name__ == "__main__":
    main()
